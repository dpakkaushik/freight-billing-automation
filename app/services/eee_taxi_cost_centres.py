"""Vehicle -> Tally cost centre master for EEE-Taxi invoices.

In Tally every income line of an invoice is tagged with a cost centre: the car
that earned the money (e.g. ``HR55AT5482(TIGOR-EV)``). An imported invoice is
only accepted when the tag already exists in Tally under exactly that name,
and the names are not uniform (some carry the model in brackets, some have
stray spaces), so the name cannot be derived from the car number.

This master keeps one row per car: normalised vehicle number -> exact Tally
name. It is filled from the "List of Cost Centres" Excel exported from Tally
(Gateway of Tally -> Chart of Accounts -> Cost Centres -> Export) and edited on
the EEE-Taxi -> Masters page behind the same edit password as the rate card.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass
from datetime import datetime
from typing import Optional

from openpyxl import load_workbook
from pydantic import BaseModel, Field, field_validator, model_validator
from sqlalchemy import delete, select
from sqlalchemy.orm import Session

from app.models import EeeTaxiCostCentre

MAX_ROWS = 2_000

# Indian registration numbers as they appear in Tally: DL1NA383, HR55AT5482, TS07UH4046.
_PLATE = re.compile(r"^[A-Z]{2}\d{1,2}[A-Z]{0,3}\d{1,4}$")
_CATEGORY_ROW = re.compile(r"cost category$", re.IGNORECASE)          # "Primary Cost Category"
_FOOTER_ROW = re.compile(r"^\d+\s+cost categor", re.IGNORECASE)       # "1 Cost Categories"


@dataclass(frozen=True)
class CostCentre:
    vehicle_no: str     # normalised, e.g. "HR55AT5482"
    cost_centre: str    # exact Tally name, e.g. "HR55AT5482(TIGOR-EV)"


@dataclass(frozen=True)
class PickedName:
    """A car that has more than one cost centre in Tally."""
    vehicle_no: str
    chosen: str
    others: tuple[str, ...]


@dataclass(frozen=True)
class TallyImport:
    rows: tuple[CostCentre, ...]
    not_cars: tuple[str, ...]        # e.g. "Head Office", "GGN"
    multi_car: tuple[str, ...]       # one name covering two cars, e.g. "A & B(TIGOR-EV)"
    picked: tuple[PickedName, ...]


def normalize_vehicle_no(value: str) -> str:
    """'hr 55 at-5482' -> 'HR55AT5482' (letters and digits only, upper case)."""
    return re.sub(r"[^A-Z0-9]", "", (value or "").upper())


def vehicle_no_from_cost_centre(name: str) -> Optional[str]:
    """Car number a Tally cost-centre name stands for, or None if it is not one car.

    'HR55AS0324 (TIGOR- EV)' -> 'HR55AS0324'; 'DL1NA 3197' -> 'DL1NA3197';
    'Head Office' -> None; 'DL52GD3187 & DL52GD3161(TIGOR-EV)' -> None.
    """
    if "&" in name:
        return None
    plate = normalize_vehicle_no(name.split("(", 1)[0])
    return plate if _PLATE.match(plate) else None


# ── Tally Excel import ────────────────────────────────────────────────────────

def _first_column(data: bytes) -> list[str]:
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises several unrelated types for a bad file
        raise ValueError("This is not an Excel (.xlsx) file. Export the cost centre list from Tally as Excel.") from exc
    try:
        values = (row[0] for row in wb.worksheets[0].iter_rows(max_col=1, values_only=True) if row)
        return [str(v).strip() for v in values if v is not None and str(v).strip()]
    finally:
        wb.close()


def _list_entries(cells: list[str]) -> list[str]:
    """Cost-centre names from a Tally list report, without its title/address/footer lines.

    Tally prints the company header first, then each cost category followed by its
    cost centres, then a "N Cost Categories and M Cost Centre(s)" footer. A sheet
    without category rows is taken to be a plain list of names.
    """
    if not any(_CATEGORY_ROW.search(c) for c in cells):
        return cells
    entries: list[str] = []
    in_list = False
    for cell in cells:
        if _FOOTER_ROW.search(cell):
            break
        if _CATEGORY_ROW.search(cell):
            in_list = True
        elif in_list:
            entries.append(cell)
    return entries


def _choose(names: list[str]) -> str:
    # A car entered twice in Tally ("DL52GD5203" and "DL52GD5203(TIGOR-EV)"):
    # the invoices in Tally use the one carrying the model.
    return next((n for n in names if "(" in n), names[0])


def parse_tally_cost_centres(data: bytes) -> TallyImport:
    """Read the cost centre list exported from Tally. Nothing is saved."""
    names = list(dict.fromkeys(_list_entries(_first_column(data))))
    by_vehicle: dict[str, list[str]] = {}
    not_cars: list[str] = []
    multi_car: list[str] = []
    for name in names:
        plate = vehicle_no_from_cost_centre(name)
        if plate is not None:
            by_vehicle.setdefault(plate, []).append(name)
        elif "&" in name:
            multi_car.append(name)
        else:
            not_cars.append(name)
    if not by_vehicle:
        raise ValueError("No car numbers found in this file. In Tally open Chart of Accounts -> Cost Centres, "
                         "export it to Excel and upload that file.")

    rows = tuple(CostCentre(plate, _choose(found)) for plate, found in sorted(by_vehicle.items()))
    picked = tuple(
        PickedName(plate, _choose(found), tuple(n for n in found if n != _choose(found)))
        for plate, found in sorted(by_vehicle.items())
        if len(found) > 1
    )
    return TallyImport(rows=rows, not_cars=tuple(not_cars), multi_car=tuple(multi_car), picked=picked)


# ── Validated input ───────────────────────────────────────────────────────────

class CostCentreIn(BaseModel):
    vehicle_no: str = Field(min_length=1, max_length=20)
    cost_centre: str = Field(min_length=1, max_length=100)

    @field_validator("vehicle_no")
    @classmethod
    def _vehicle(cls, v: str) -> str:
        plate = normalize_vehicle_no(v)
        if not plate:
            raise ValueError("Vehicle number cannot be blank.")
        return plate

    @field_validator("cost_centre")
    @classmethod
    def _name(cls, v: str) -> str:
        # Only the ends are trimmed: inner spacing is part of the Tally name.
        v = v.strip()
        if not v:
            raise ValueError("Cost centre name cannot be blank.")
        return v


class CostCentresIn(BaseModel):
    """The whole cost centre list as saved from the Masters page."""
    rows: list[CostCentreIn] = Field(default_factory=list, max_length=MAX_ROWS)

    @model_validator(mode="after")
    def _unique(self) -> "CostCentresIn":
        seen: set[str] = set()
        for r in self.rows:
            if r.vehicle_no in seen:
                raise ValueError(f"Vehicle {r.vehicle_no} is listed twice.")
            seen.add(r.vehicle_no)
        return self

    def to_rows(self) -> tuple[CostCentre, ...]:
        return tuple(sorted((CostCentre(r.vehicle_no, r.cost_centre) for r in self.rows),
                            key=lambda c: c.vehicle_no))


# ── Persistence ───────────────────────────────────────────────────────────────

@dataclass(frozen=True)
class CostCentreState:
    rows: tuple[CostCentre, ...]
    updated_by: Optional[str]
    updated_at: Optional[datetime]


def get_cost_centre_state(db: Session) -> CostCentreState:
    records = db.scalars(select(EeeTaxiCostCentre).order_by(EeeTaxiCostCentre.vehicle_no)).all()
    latest = max(records, key=lambda r: r.updated_at, default=None)
    return CostCentreState(
        rows=tuple(CostCentre(r.vehicle_no, r.cost_centre) for r in records),
        updated_by=latest.updated_by if latest else None,
        updated_at=latest.updated_at if latest else None,
    )


def save_cost_centres(db: Session, rows: tuple[CostCentre, ...], updated_by: str) -> None:
    """Replace the whole list in one transaction."""
    now = datetime.utcnow()
    db.execute(delete(EeeTaxiCostCentre))
    db.add_all([
        EeeTaxiCostCentre(vehicle_no=r.vehicle_no, cost_centre=r.cost_centre, updated_by=updated_by, updated_at=now)
        for r in rows
    ])
    db.commit()
